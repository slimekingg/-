import webbrowser
import openpyxl

for ii in range(1,6):
    wb = openpyxl.load_workbook('点名情况课程'+str(ii)+'.xlsx')
    ws = wb['课程'+str(ii)]

    a = []
    for i in range(1,91):
        a.append((ws.cell(i+1,22).value*10,i))
    a.sort()

    x=0
    y=0
    k=30
    mi=8
    for i in range(1,21):
        if i==1 : k=20
        if i==2 : k=13
        if i==3 : k=11
        if i==4 : k=10
        if i>4 : k=mi
        print('第'+str(ii)+'门课第'+str(i)+'次课点:')
        for j in range(k):
            print(ws.cell(a[j][1]+1,1).value,end=" ")
            y+=1
            if ws.cell(a[j][1]+1,i+1).value==0 :
                a[j]=(a[j][0]-50,a[j][1])
                x+=1
            if i>4 :
                if a[j][0]>-100 :
                    mi-=1
        a.sort()
        print(" ")
    print("E值为",x/y)
